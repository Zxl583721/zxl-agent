from pathlib import Path

SUPPORTED_EXTENSIONS = {".txt", ".pdf", ".docx", ".pptx"}


def load_documents(data_dir: str | Path) -> list[dict]:
    """Load supported knowledge base files from the data directory."""
    page_documents = load_document_pages(data_dir)
    grouped_documents: dict[str, list[str]] = {}

    for page in page_documents:
        grouped_documents.setdefault(page["source"], []).append(page["text"])

    return [
        {
            "source": source,
            "content": clean_text("\n\n".join(texts)),
        }
        for source, texts in grouped_documents.items()
        if clean_text("\n\n".join(texts))
    ]


def load_document_pages(data_dir: str | Path) -> list[dict]:
    """Load supported knowledge base files while preserving page or slide metadata."""
    data_path = Path(data_dir)
    if not data_path.exists():
        return []

    pages = []
    for file_path in sorted(data_path.iterdir()):
        if not file_path.is_file() or file_path.suffix.lower() not in SUPPORTED_EXTENSIONS:
            continue

        for page in load_file_pages(file_path):
            text = clean_text(page.get("text", ""))
            if text:
                pages.append(
                    {
                        "source": str(file_path),
                        "filename": file_path.name,
                        "page": page.get("page"),
                        "text": text,
                    }
                )

    return pages


def load_txt_documents(data_dir: str | Path) -> list[dict]:
    """Backward-compatible alias for older callers."""
    return load_documents(data_dir)


def load_file_content(file_path: Path) -> str:
    """Extract plain text from one supported knowledge base file."""
    return "\n\n".join(page["text"] for page in load_file_pages(file_path)).strip()


def load_file_pages(file_path: Path) -> list[dict]:
    """Extract plain text pages or slides from one supported knowledge base file."""
    suffix = file_path.suffix.lower()

    if suffix == ".txt":
        content = load_txt_content(file_path)
        return [{"page": 1, "text": content}] if content else []
    if suffix == ".pdf":
        return load_pdf_pages(file_path)
    if suffix == ".docx":
        content = load_docx_content(file_path)
        return [{"page": 1, "text": content}] if content else []
    if suffix == ".pptx":
        return load_pptx_pages(file_path)

    return []


def clean_text(text: str) -> str:
    """Remove invalid Unicode code points that cannot be sent as JSON."""
    if not text:
        return ""

    return text.encode("utf-8", errors="ignore").decode("utf-8").strip()


def load_txt_content(file_path: Path) -> str:
    try:
        return file_path.read_text(encoding="utf-8").strip()
    except UnicodeDecodeError:
        return file_path.read_text(encoding="gbk", errors="ignore").strip()


def load_pdf_content(file_path: Path) -> str:
    return "\n\n".join(page["text"] for page in load_pdf_pages(file_path)).strip()


def load_pdf_pages(file_path: Path) -> list[dict]:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise RuntimeError("读取 PDF 需要先安装 pypdf，请运行 pip install -r requirements.txt。") from exc

    reader = PdfReader(str(file_path))
    pages = []
    for page_number, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if text:
            pages.append({"page": page_number, "text": f"第 {page_number} 页\n{text}"})
    return pages


def load_docx_content(file_path: Path) -> str:
    try:
        from docx import Document as DocxDocument
    except ImportError as exc:
        raise RuntimeError("读取 Word 需要先安装 python-docx，请运行 pip install -r requirements.txt。") from exc

    document = DocxDocument(str(file_path))
    paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs]
    tables = []

    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                tables.append(" | ".join(cells))

    return "\n".join([text for text in paragraphs + tables if text]).strip()


def load_pptx_content(file_path: Path) -> str:
    return "\n\n".join(page["text"] for page in load_pptx_pages(file_path)).strip()


def load_pptx_pages(file_path: Path) -> list[dict]:
    try:
        from pptx import Presentation
    except ImportError as exc:
        raise RuntimeError("读取 PPT 需要先安装 python-pptx，请运行 pip install -r requirements.txt。") from exc

    presentation = Presentation(str(file_path))
    slides = []

    for slide_number, slide in enumerate(presentation.slides, start=1):
        texts = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                texts.append(shape.text.strip())

            if not getattr(shape, "has_table", False):
                continue

            for row in shape.table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    texts.append(" | ".join(cells))

        if texts:
            slides.append({"page": slide_number, "text": f"第 {slide_number} 页\n" + "\n".join(texts)})

    return slides
